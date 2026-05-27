import os
import logging
from pathlib import Path

import litellm

logger = logging.getLogger("gerente_ia.tools.rag")

CHROMA_DIR = Path(__file__).parent.parent / "chroma_db"
DOCS_DIR = Path(__file__).parent.parent / "documents"

CHUNK_SIZE = 800
CHUNK_OVERLAP = 100

try:
    import chromadb
    CHROMA_AVAILABLE = True
except ImportError:
    CHROMA_AVAILABLE = False
    logger.warning("chromadb no instalado — RAG desactivado")


def _get_collection(connection_id: str):
    CHROMA_DIR.mkdir(parents=True, exist_ok=True)
    client = chromadb.PersistentClient(path=str(CHROMA_DIR))
    name = f"kb_{connection_id.replace('-', '_')}"
    return client.get_or_create_collection(name=name, metadata={"hnsw:space": "cosine"})


def _embed(text: str) -> list[float]:
    response = litellm.embedding(
        model="ollama/nomic-embed-text",
        input=[text],
        api_base=os.getenv("OLLAMA_BASE_URL", "http://localhost:11434"),
    )
    return response.data[0]["embedding"]


def _chunk_text(text: str) -> list[str]:
    chunks = []
    start = 0
    while start < len(text):
        end = start + CHUNK_SIZE
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        start = end - CHUNK_OVERLAP
    return chunks


def _extract_text(file_path: Path, filename: str) -> str:
    ext = filename.rsplit(".", 1)[-1].lower()
    try:
        if ext == "pdf":
            import pdfplumber
            with pdfplumber.open(file_path) as pdf:
                return "\n".join(page.extract_text() or "" for page in pdf.pages)

        elif ext in ("doc", "docx"):
            from docx import Document
            doc = Document(file_path)
            return "\n".join(p.text for p in doc.paragraphs)

        elif ext in ("xlsx", "xls"):
            import openpyxl
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            lines = []
            for sheet in wb.worksheets:
                lines.append(f"[Hoja: {sheet.title}]")
                for row in sheet.iter_rows(values_only=True):
                    lines.append("\t".join(str(c) if c is not None else "" for c in row))
            return "\n".join(lines)

        elif ext == "csv":
            return file_path.read_text(encoding="utf-8", errors="replace")

        else:
            return file_path.read_text(encoding="utf-8", errors="replace")

    except Exception as e:
        logger.warning(f"Error extrayendo texto de {filename}: {e}")
        return ""


def ingest_file(connection_id: str, file_path: Path, filename: str) -> dict:
    if not CHROMA_AVAILABLE:
        return {"success": False, "error": "chromadb no instalado"}

    text = _extract_text(file_path, filename)
    if not text.strip():
        return {"success": False, "error": "No se pudo extraer texto del archivo"}

    chunks = _chunk_text(text)
    if not chunks:
        return {"success": False, "error": "El archivo no contiene texto procesable"}

    collection = _get_collection(connection_id)

    # Eliminar versión anterior del mismo archivo si existe
    try:
        existing = collection.get(where={"filename": filename})
        if existing["ids"]:
            collection.delete(ids=existing["ids"])
    except Exception:
        pass

    ids, embeddings, documents, metadatas = [], [], [], []
    for i, chunk in enumerate(chunks):
        try:
            emb = _embed(chunk)
        except Exception as e:
            logger.error(f"Error embedding chunk {i} de {filename}: {e}")
            return {"success": False, "error": f"Error al generar embeddings: {e}"}
        ids.append(f"{connection_id}__{filename}__{i}")
        embeddings.append(emb)
        documents.append(chunk)
        metadatas.append({"filename": filename, "chunk_index": i, "connection_id": connection_id})

    collection.add(ids=ids, embeddings=embeddings, documents=documents, metadatas=metadatas)
    logger.info(f"[RAG] Ingestado '{filename}' → {len(chunks)} chunks en colección '{connection_id}'")
    return {"success": True, "chunks": len(chunks), "filename": filename}


def delete_document(connection_id: str, filename: str) -> None:
    if not CHROMA_AVAILABLE:
        return
    try:
        collection = _get_collection(connection_id)
        existing = collection.get(where={"filename": filename})
        if existing["ids"]:
            collection.delete(ids=existing["ids"])
        logger.info(f"[RAG] Eliminado '{filename}' de colección '{connection_id}'")
    except Exception as e:
        logger.warning(f"[RAG] Error eliminando '{filename}': {e}")


def list_documents(connection_id: str) -> list[dict]:
    if not CHROMA_AVAILABLE:
        return []
    try:
        collection = _get_collection(connection_id)
        result = collection.get(include=["metadatas"])
        seen: dict[str, dict] = {}
        for meta in result["metadatas"]:
            fn = meta["filename"]
            if fn not in seen:
                seen[fn] = {"filename": fn, "chunks": 0}
            seen[fn]["chunks"] += 1
        return list(seen.values())
    except Exception as e:
        logger.warning(f"[RAG] Error listando documentos de '{connection_id}': {e}")
        return []


def search(connection_id: str, query: str, n_results: int = 5) -> dict:
    if not CHROMA_AVAILABLE:
        return {"success": False, "error": "chromadb no instalado", "row_count": 0}
    try:
        collection = _get_collection(connection_id)
        total = collection.count()
        if total == 0:
            return {
                "success": False,
                "error": "La base de conocimiento está vacía. No hay documentos indexados.",
                "row_count": 0,
            }

        query_emb = _embed(query)
        results = collection.query(
            query_embeddings=[query_emb],
            n_results=min(n_results, total),
            include=["documents", "metadatas", "distances"],
        )

        chunks_found = []
        for doc, meta, dist in zip(
            results["documents"][0],
            results["metadatas"][0],
            results["distances"][0],
        ):
            chunks_found.append({
                "content": doc,
                "filename": meta["filename"],
                "relevance": round(1 - dist, 3),
            })

        return {
            "success": True,
            "query": query,
            "results": chunks_found,
            "row_count": len(chunks_found),
        }
    except Exception as e:
        logger.error(f"[RAG] Error en búsqueda: {e}")
        return {"success": False, "error": str(e), "row_count": 0}


def build_tool_definition(connection_id: str, connection_name: str, doc_count: int = 0) -> dict:
    doc_info = f" Contiene {doc_count} documento(s) indexados." if doc_count > 0 else " Aún sin documentos."
    return {
        "type": "function",
        "function": {
            "name": f"search_knowledge_base_{connection_id.replace('-', '_')}",
            "description": (
                f"Busca información en la base de conocimiento '{connection_name}'.{doc_info} "
                "Úsala cuando necesites información sobre políticas, procedimientos, reportes internos, "
                "manuales, contratos o cualquier documento de la empresa. "
                "Formula la búsqueda como una pregunta o descripción detallada de lo que necesitas."
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "query": {
                        "type": "string",
                        "description": "Pregunta o descripción detallada de la información que se busca",
                    }
                },
                "required": ["query"],
            },
        },
    }
